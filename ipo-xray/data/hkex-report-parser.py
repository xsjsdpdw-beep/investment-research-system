import json
import sys
from datetime import date, datetime

from openpyxl import load_workbook


def iso(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return None


def clean(value):
    return " ".join(str(value or "").replace("\n", " ").split())


def number(value):
    return value if isinstance(value, (int, float)) else None


rows = []
for file_path in sys.argv[1:]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    pending = None
    for values in sheet.iter_rows(min_row=3, values_only=True):
        sequence, code, name = values[0], values[1], values[2]
        note = clean(values[10] if len(values) > 10 else "")
        if isinstance(sequence, (int, float)) and code and name:
            pending = {
                "sequence": int(sequence),
                "code": str(code).zfill(5),
                "name": clean(name),
                "prospectusDate": iso(values[3]),
                "listingDate": iso(values[4]),
                "sponsor": clean(values[5]),
                "accountant": clean(values[6]),
                "valuer": clean(values[7]),
                "offerShares": number(values[8]),
                "issuePrice": number(values[9]),
                "fundsRaised": None,
            }
            rows.append(pending)
            continue
        if pending and str(code or "").strip() == '"' and note.lower().startswith("(b"):
            pending["fundsRaised"] = number(values[8])

sys.stdout.write(json.dumps(rows, ensure_ascii=False))
